"""기획서 목차 쪽 번호 (D-0246).

**LibreOffice를 머리 없이 돌리면 Word 목차 필드를 갱신하지 않는다** — 실측으로 확인했고
자리표시자가 그대로 인쇄됐다. 그래서 굽고 · 세고 · 붙이고 · 다시 굽는다.

굽는 것은 여기서 안 돌린다(LibreOffice가 필요하다). **세는 자리와 붙이는 자리**를 본다.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from hathor.shared.config.paths import repo_root

_spec = importlib.util.spec_from_file_location(
    "bake_proposal", repo_root() / "tools" / "bake_proposal.py"
)
assert _spec is not None and _spec.loader is not None
BAKER = importlib.util.module_from_spec(_spec)
sys.modules["bake_proposal"] = BAKER
_spec.loader.exec_module(BAKER)

TITLES = ["Part I. 프로젝트 제안서", "1. 프로젝트 개요", "12. 산출물 목록"]


def test_목차_쪽이_아니라_본문_쪽을_센다() -> None:
    """**목차에도 같은 글자가 있다.** 거기서 세면 전부 목차 쪽이 된다."""
    pages = [
        "표지",
        "목차\nPart I. 프로젝트 제안서\n1. 프로젝트 개요\n12. 산출물 목록",
        "Part I. 프로젝트 제안서\n1. 프로젝트 개요\n본문",
        "본문",
        "12. 산출물 목록\n본문",
    ]

    assert BAKER.locate(pages, TITLES) == {
        "Part I. 프로젝트 제안서": 3,
        "1. 프로젝트 개요": 3,
        "12. 산출물 목록": 5,
    }


def test_줄이_갈려도_찾는다() -> None:
    """PDF는 줄바꿈과 공백을 제 마음대로 넣는다. **띄어쓰기를 지우고 본다.**"""
    pages = ["목차\n1. 프로젝트 개요", "1.  프로젝트\n개요\n본문"]

    assert BAKER.locate(pages, ["1. 프로젝트 개요"]) == {"1. 프로젝트 개요": 2}


def test_본문에_없는_항목은_안_적는다() -> None:
    """못 찾은 것을 0쪽으로 적으면 **거짓말이 인쇄된다.** 비워 두고 호출자가 실패한다."""
    assert BAKER.locate(["목차\n없는 제목"], ["없는 제목"]) == {}


def test_목차_줄에_점선과_쪽을_붙인다(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "before.docx"
    document = Document()
    document.add_paragraph("목차")
    for title in TITLES:
        document.add_paragraph(title)
    document.add_paragraph("여기부터는 본문이다")
    document.save(str(source))

    target = tmp_path / "after.docx"
    written = BAKER.stamp(source, target, {TITLES[0]: 4, TITLES[1]: 4, TITLES[2]: 66})

    assert written == 3
    rows = [p.text for p in Document(str(target)).paragraphs]
    assert rows[1].endswith("\t4") and rows[3].endswith("\t66")
    assert rows[4] == "여기부터는 본문이다", "본문은 안 건드린다"


def test_목차_뒤가_끝나면_멈춘다(tmp_path: Path) -> None:
    """본문에 같은 글자가 또 나와도 **목차 블록을 벗어나면 안 붙인다.**"""
    from docx import Document

    source = tmp_path / "before.docx"
    document = Document()
    document.add_paragraph("목차")
    document.add_paragraph(TITLES[0])
    document.add_paragraph("본문 시작")
    document.add_paragraph(TITLES[0])
    document.save(str(source))

    target = tmp_path / "after.docx"
    assert BAKER.stamp(source, target, {TITLES[0]: 4}) == 1
    rows = [p.text for p in Document(str(target)).paragraphs]
    assert rows[1].endswith("\t4") and rows[3] == TITLES[0]
