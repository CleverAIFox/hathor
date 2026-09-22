#!/usr/bin/env python3
"""`MASTER.md` Part I ~ III와 그림으로 기획서 docx를 빌드한다 (D-0221).

### 왜 빌드하나

D-0220은 기획서를 **따로 손으로** 썼다. 8쪽 요약이 나왔고 `MASTER.md`의 3부(제안서 ·
요구사항 분석서 · 상세 설계서)와 내용이 갈렸다 — 같은 것을 두 곳에 쓰면 한쪽만 고쳐진다.
fire-lane 기획서가 62쪽 · 그림 24장 · 표 56개인 것은 **설계 문서 전체가 기획서**이기 때문이다.

그래서 기획서는 **빌드 산출물**이다. 정본은 `MASTER.md` 하나이고 이 도구가 3부를 docx로
옮기며, 그림은 `render_figures.py` · `render_charts.py`가 **같은 정본에서** 그린다.

### 낡으면 빨개진다

빌드가 정본 구간의 지문을 docx 속성에 적는다. `docx_check.py`가 CI에서 지금 지문과 맞댄다 —
Part I ~ III를 고치고 다시 빌드하지 않으면 **검사가 멈춘다.**

### 필요한 것

`uv sync --group docs`(matplotlib · python-docx · pandoc 동봉) · graphviz(`dot`) · 한글 글꼴.

    make proposal
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from proposal_source import FINGERPRINT, ROOT, SourceError, fingerprint, source

OUT = "docs/proposal.docx"
FONT = "맑은 고딕"
MONO = "Consolas"
NAVY = "1F4E79"
RUST = "A33A2B"
HEAD_FILL = "E8EEF5"
PAGE_WIDTH_IN = 6.3
PAGE_HEIGHT_IN = 8.3

# (제목 줄, 그림, 캡션, 자리) — 자리: replace = 절의 첫 코드 블록을 바꾼다
#                                     intro = 절의 첫 문단 뒤 · end = 절 끝
FIGURES: list[tuple[str, str, str, str]] = [
    (
        "### 목적 및 핵심가치",
        "concept",
        "서비스 개념도 — 음원은 기기를 떠나지 않고 벡터만 기획으로 간다",
        "end",
    ),
    ("### □ 추진 일정", "gantt", "추진 일정과 단계별 상태", "end"),
    (
        "## □ REQ-HATHOR 요구사항 체계도",
        "requirement_tree",
        "요구사항 체계도 — 12개 영역",
        "replace",
    ),
    ("## □ REQ-HATHOR 요구사항 체계도", "use_cases", "유스케이스 — 사용자 · 개발자", "end"),
    (
        "## □ 세부 기능 요구사항 정의서",
        "requirement_status",
        "요구사항 상태 집계 — 영역별",
        "intro",
    ),
    (
        "### □ 전체 파이프라인",
        "pipeline",
        "전체 파이프라인 — 기기 분석 · 기획 · 렌더 · 관문",
        "replace",
    ),
    ("### □ 계층 구조 (GR-2)", "layers", "계층 구조와 import 계약 5종", "end"),
    ("### □ 테이블 명세 (주요) — P4 설계 · 미구현", "erd", "ERD — P4 설계", "intro"),
    ("### □ ID3 프레임 실측 (1004곡 전수)", "id3_frames", "ID3 프레임 보유율 (1004곡 전수)", "end"),
    ("### □ 아티스트 표기 실측", "artist_notation", "아티스트 표기 구성", "end"),
    ("### □ 처리 순서", "ingest_flow", "인제스트 처리 순서", "replace"),
    ("### □ 델타 이벤트", "delta_states", "델타 판정 — 재스캔 이벤트 넷", "end"),
    (
        "### □ MusicBrainz 조회 개선 이력 (D-0019)",
        "musicbrainz",
        "MusicBrainz 곡 조회 성공률 개선",
        "end",
    ),
    ("## 4. 4축 특징 설계", "axes", "4축과 모델 — 스템 분리가 편곡 · 가창의 전제다", "end"),
    ("### □ 취향 신호 3층", "taste_layers", "취향 신호 3층", "end"),
    (
        "### □ 기획과 렌더를 나눈다",
        "plan_render",
        "hathor가 기획하고 엔진이 소리를 입힌다",
        "replace",
    ),
    (
        "### □ 엔진 선정 — 서류 관문 (D-0217)",
        "engine_funnel",
        "엔진 서류 관문 — 일곱에서 둘로",
        "end",
    ),
    (
        "### □ 실측 관문 — 착수 전에 적었다 (D-0217)",
        "engine_gates",
        "엔진 실측 관문 G0 ~ G3",
        "end",
    ),
    (
        "### □ G0 실측 — 기기가 떨어졌다 (D-0218)",
        "g0_runs",
        "G0 실측 — fp16 NaN · fp32 메모리 부족",
        "end",
    ),
    ("### □ 청취 판정 이력", "listening", "청취 판정 이력", "end"),
    ("### □ 청취 판정 이력", "voice_leading", "성부 진행 전후 (D-0137)", "end"),
    ("### □ 음색 누설 관문 (G2)", "leak_gate", "음색 누설 관문 G2", "replace"),
    ("### □ 지표", "retrieval", "검색 평가 — 무작위 대비", "end"),
    ("### □ 레이어 곡선 — 검색축 (D-0027)", "layer_curve", "레이어 곡선 — 검색축", "end"),
    ("### □ 층별 화성 — 화성축 (D-0181)", "harmony_layers", "층별 화성 — 화성축", "end"),
    ("### □ 패치 파이프", "patch_pipe", "패치 파이프", "replace"),
    ("### □ 검사 체계", "gate_parity", "make check · CI · 커밋 훅 대조 (실물에서 읽었다)", "end"),
    ("### □ 문서 체계", "documents", "문서 체계 — 기획서는 MASTER에서 빌드한다", "end"),
]


# ------------------------------------------------------------------ 마크다운


def _span(lines: list[str], heading: str) -> tuple[int, int]:
    for index, line in enumerate(lines):
        if line.strip() == heading:
            level = len(line) - len(line.lstrip("#"))
            for later in range(index + 1, len(lines)):
                match = re.match(r"^(#+) ", lines[later])
                if match and len(match.group(1)) <= level:
                    return index, later
            return index, len(lines)
    raise SourceError(
        f"그림 자리 `{heading}`이 `MASTER.md`에 없다 — 제목을 바꿨으면 FIGURES도 고친다"
    )


def _image(path: Path, caption: str) -> list[str]:
    from PIL import Image  # matplotlib이 끌고 온다

    with Image.open(path) as picture:
        width, height = picture.size
    inches = min(PAGE_WIDTH_IN, PAGE_HEIGHT_IN * width / height, width / 150)
    return ["", f"![{caption}]({path.resolve().as_posix()}){{width={inches:.2f}in}}", ""]


def place_figures(body: str, figures: dict[str, Path]) -> str:
    lines = body.splitlines()
    for heading, name, caption, where in FIGURES:
        start, end = _span(lines, heading)
        image = _image(figures[name], caption)
        if where == "replace":
            fence = next((i for i in range(start, end) if lines[i].startswith("```")), None)
            if fence is None:
                raise SourceError(f"`{heading}`에 바꿀 코드 블록이 없다")
            close = next(i for i in range(fence + 1, end) if lines[i].startswith("```"))
            lines[fence : close + 1] = image
        elif where == "intro":
            first = next(i for i in range(start + 1, end) if lines[i].strip())
            blank = next((i for i in range(first, end) if not lines[i].strip()), end)
            lines[blank:blank] = image
        else:
            while end > start and not lines[end - 1].strip():
                end -= 1
            lines[end:end] = image
    return "\n".join(lines) + "\n"


def split_cover(text: str) -> tuple[list[str], str]:
    """표지 문단들과 본문. 마크다운은 문단 안에서 줄을 접으므로 **빈 줄로 문단을 가른다.**"""
    head, _, body = text.partition("# Part I.")
    paragraphs = [
        " ".join(line.strip() for line in block.splitlines())
        for block in re.split(r"\n\s*\n", head)
        if block.strip() and block.strip() != "---"
    ]
    return paragraphs, "# Part I." + body


def clean(markdown: str) -> str:
    markdown = re.sub(r"<!--.*?-->", "", markdown, flags=re.S)
    return markdown.replace("\n---\n", "\n")


# ------------------------------------------------------------------ pandoc


def _pandoc() -> str:
    try:
        import pypandoc

        return str(pypandoc.get_pandoc_path())
    except (ImportError, OSError):
        found = shutil.which("pandoc")
        if not found:
            raise SourceError("pandoc이 없다. `cd core && uv sync --group docs`") from None
        return found


def reference_doc(work: Path, pandoc: str) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    target = work / "reference.docx"
    data = subprocess.run(
        [pandoc, "--print-default-data-file", "reference.docx"], capture_output=True, check=True
    ).stdout
    target.write_bytes(data)
    document = Document(str(target))
    styles = document.styles

    def font(style: Any, size: float, bold: bool = False, color: str | None = None) -> None:
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor.from_string(color)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if fonts is None:
            fonts = rpr.makeelement(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts", {}
            )
            rpr.append(fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(
                f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{key}", FONT
            )

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            font(styles[name], 9.5)
            styles[name].paragraph_format.space_after = Pt(4)
            styles[name].paragraph_format.line_spacing = 1.25
    for name, size in (
        ("Heading 1", 20),
        ("Heading 2", 14),
        ("Heading 3", 11.5),
        ("Heading 4", 10.5),
    ):
        font(styles[name], size, True, NAVY)
        styles[name].paragraph_format.space_before = Pt(14 if size > 11 else 10)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.space_after = Pt(18)
    if "Image Caption" in [s.name for s in styles]:
        font(styles["Image Caption"], 8.5, False, "555555")
        styles["Image Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Source Code" in [s.name for s in styles]:
        styles["Source Code"].font.size = Pt(7.5)
    document.save(str(target))
    return target


# ------------------------------------------------------------------ 다듬기


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "BFC8D3")
        borders.append(element)
    tbl_pr.append(borders)


TBL_PR_ORDER = (
    "tblStyle",
    "tblpPr",
    "tblOverlap",
    "bidiVisual",
    "tblStyleRowBandSize",
    "tblStyleColBandSize",
    "tblW",
    "jc",
    "tblCellSpacing",
    "tblInd",
    "tblBorders",
    "shd",
    "tblLayout",
    "tblCellMar",
    "tblLook",
    "tblCaption",
    "tblDescription",
)
"""`w:tblPr` 자식의 스키마 순서. 어기면 Word가 «복구»를 묻는다."""


def _order(tbl_pr: Any) -> None:
    def rank(element: Any) -> int:
        name = element.tag.rsplit("}", 1)[-1]
        return TBL_PR_ORDER.index(name) if name in TBL_PR_ORDER else len(TBL_PR_ORDER)

    children = sorted(tbl_pr, key=rank)
    for child in list(tbl_pr):
        tbl_pr.remove(child)
    for child in children:
        tbl_pr.append(child)


def _full_width(table: Any, inches: list[float]) -> None:
    """표 폭을 본문 폭에 맞춘다. pandoc은 격자 열 폭을 좁게 박아 LibreOffice가 표를 줄인다."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(sum(inches) * 1440)))
    width.set(qn("w:type"), "dxa")
    tbl_pr.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    _order(tbl_pr)
    grid = table._tbl.tblGrid
    for column, size in zip(grid.findall(qn("w:gridCol")), inches, strict=False):
        column.set(qn("w:w"), str(int(size * 1440)))


def column_widths(table: Any, columns: int) -> list[float]:
    """열마다 **가장 긴 낱말이 안 끊기는 폭**을 먼저 주고, 남는 폭을 글 양에 비례해 나눈다.

    글 양만으로 나누면 «완료» 같은 짧은 열이 한 글자 폭으로 줄어 세로로 찢어진다.
    """
    floors, loads = [0.0] * columns, [0.0] * columns
    for row in table.rows:
        for index, cell in enumerate(row.cells[:columns]):
            text = cell.text.strip()
            longest = max((len(word) for word in text.split()), default=1)
            floors[index] = max(floors[index], min(0.12 + 0.105 * longest, 1.3))
            loads[index] += len(text)
    spare = max(PAGE_WIDTH_IN - sum(floors), 0.0)
    total = sum(load**0.9 for load in loads) or 1.0
    widths = [
        floor + spare * (load**0.9) / total for floor, load in zip(floors, loads, strict=True)
    ]
    scale = PAGE_WIDTH_IN / sum(widths)
    return [width * scale for width in widths]


def polish_tables(document: Any) -> int:
    from docx.shared import Inches, Pt, RGBColor

    for table in document.tables:
        _borders(table)
        columns = len(table.columns)
        widths = column_widths(table, columns)
        _full_width(table, widths)
        for row_index, row in enumerate(table.rows):
            for index, cell in enumerate(row.cells[:columns]):
                cell.width = Inches(widths[index])
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.1
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
                        if row_index == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(NAVY)
                if row_index == 0:
                    _shade(cell, HEAD_FILL)
    return len(document.tables)


def number_figures(document: Any) -> int:
    count = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Image Caption" and paragraph.runs:
            count += 1
            first = paragraph.runs[0]
            first.text = f"[그림 {count}] " + first.text
    return count


def cover(document: Any, lines: list[str], parts: list[str], stamp: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, RGBColor

    body = document.element.body
    first = body[0]
    made: list[Any] = []

    def add(
        text: str,
        size: float,
        bold: bool = False,
        color: str = "333333",
        center: bool = True,
        before: float = 0,
        after: float = 6,
    ) -> Any:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        for chunk in re.split(r"(\*\*[^*]+\*\*)", text):
            if not chunk:
                continue
            run = paragraph.add_run(chunk.strip("*") if chunk.startswith("**") else chunk)
            run.font.size = Pt(size)
            run.font.bold = bold or chunk.startswith("**")
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.name = FONT
        made.append(paragraph)
        return paragraph

    title = lines[0].lstrip("# ").strip()
    add("프로젝트 기획서", 16, True, NAVY, before=120)
    add("HATHOR", 44, True, RUST, after=4)
    add(title, 15, True, "333333", after=18)
    notes = [line for line in lines[1:] if line.startswith("**※")]
    for line in (line for line in lines[1:] if not line.startswith("**※")):
        add(line.replace("`", "").replace(" | ", "\n"), 10.5, color="555555", after=6)
    add(f"빌드 {stamp} · 정본 docs/MASTER.md Part I ~ III", 8.5, color="888888", before=30)
    breaker = document.add_paragraph()
    breaker.add_run().add_break(WD_BREAK.PAGE)
    made.append(breaker)
    add("전제", 14, True, NAVY, center=False, after=8)
    for note in notes:
        add(note.replace("`", ""), 8.5, color="333333", center=False, after=6)
    toc_break = document.add_paragraph()
    toc_break.add_run().add_break(WD_BREAK.PAGE)
    made.append(toc_break)
    add("목차", 14, True, NAVY, center=False, after=8)
    for part in parts:
        is_part = part.startswith("Part")
        add(
            part,
            10.5 if is_part else 9,
            is_part,
            NAVY if is_part else "333333",
            center=False,
            before=6 if is_part else 0,
            after=1,
        )
    for paragraph in made:
        first.addprevious(paragraph._p)


def outline(markdown: str) -> list[str]:
    found = []
    for line in markdown.splitlines():
        if line.startswith("# Part"):
            found.append(line[2:].strip())
        elif re.match(r"^## (\d+\.|□ )", line):
            found.append("    " + line[3:].strip())
    return found


def footer(document: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.95)
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.header_distance = section.footer_distance = Inches(0.45)
    section.gutter = Inches(0)
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("HATHOR 기획서 · ")
    run.font.size = Pt(8)
    field = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.text = text
        field._r.append(element)
    field.font.size = Pt(8)


def freeze(path: Path) -> None:
    """zip 항목 시각을 고정한다 — 같은 입력이면 같은 바이트에 가깝게."""
    temporary = path.with_suffix(".tmp")
    with (
        zipfile.ZipFile(path) as source_zip,
        zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as out,
    ):
        for item in sorted(source_zip.infolist(), key=lambda entry: entry.filename):
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(info, source_zip.read(item.filename))
    temporary.replace(path)


# ------------------------------------------------------------------ 빌드


def build(out: Path, figures_dir: Path) -> dict[str, int]:
    import render_charts
    import render_figures
    from docx import Document

    figures = {**render_figures.render_all(figures_dir), **render_charts.render_all(figures_dir)}
    missing = {name for _, name, _, _ in FIGURES} - set(figures)
    if missing:
        raise SourceError(f"그리지 않은 그림을 부른다: {sorted(missing)}")
    text = source()
    cover_lines, body = split_cover(text)
    markdown = clean(place_figures(body, figures))
    pandoc = _pandoc()
    with tempfile.TemporaryDirectory() as work_dir:
        work = Path(work_dir)
        (work / "body.md").write_text(markdown, encoding="utf-8")
        reference = reference_doc(work, pandoc)
        raw = work / "raw.docx"
        subprocess.run(
            [
                pandoc,
                str(work / "body.md"),
                "-f",
                "markdown+pipe_tables-yaml_metadata_block",
                "-t",
                "docx",
                "--reference-doc",
                str(reference),
                "-o",
                str(raw),
            ],
            check=True,
            cwd=ROOT,
        )
        document = Document(str(raw))
    stamp = dt.date.today().isoformat()
    tables = polish_tables(document)
    count = number_figures(document)
    cover(document, cover_lines, outline(markdown), stamp)
    footer(document)
    properties = document.core_properties
    properties.title = "HATHOR 기획서"
    properties.author = "오창준"
    properties.subject = "취향 잠재 표현 기반 종단간 AI 음악 창작 시스템"
    properties.comments = f"{FINGERPRINT} {fingerprint()}"
    fixed = dt.datetime(2026, 1, 1, tzinfo=dt.UTC)
    properties.created = properties.modified = fixed
    properties.revision = 1
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    freeze(out)
    if count != len(FIGURES):
        raise SourceError(f"그림 {len(FIGURES)}장을 넣었는데 캡션이 {count}개다")
    return {"figures": count, "tables": tables}


def main() -> int:
    parser = argparse.ArgumentParser(description="기획서 빌드 (D-0221)")
    parser.add_argument("--out", type=Path, default=ROOT / OUT)
    parser.add_argument("--figures", type=Path, default=ROOT / "var" / "proposal" / "figures")
    args = parser.parse_args()
    try:
        made = build(args.out, args.figures)
    except SourceError as error:
        print(f"기획서를 빌드할 수 없다: {error}", file=sys.stderr)
        return 2
    print(f"기획서 빌드 · 그림 {made['figures']}장 · 표 {made['tables']}개 · {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
