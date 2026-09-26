#!/usr/bin/env python3
"""기획서를 PDF로 굽고 **목차에 쪽 번호를 박는다** (D-0246).

### 왜 PDF에만 박는가

`.docx`에는 쪽이 없다. 여는 프로그램·글꼴·용지에 따라 달라지므로 **거기 숫자를 적으면
거짓말이 된다.** 쪽은 굽힌 PDF의 성질이고, 그래서 숫자도 거기서 세어 거기에만 넣는다.

### 두 판을 굽는다

1. 첫 판을 굽고 제목마다 몇 쪽인지 센다.
2. 목차 줄에 숫자를 붙인 사본을 만들어 다시 굽는다.
3. **다시 세어 같은지 본다.** 숫자는 줄 오른쪽에 붙으므로 줄 수가 안 변하고, 따라서 쪽도
   안 밀린다 — 그 전제가 깨지면(목차가 한 쪽을 넘기는 등) **실패로 끝낸다.**

`--updateFields` 같은 길은 막혀 있다. LibreOffice를 머리 없이 돌리면 Word의 목차 필드를
**갱신하지 않고** 자리표시자를 그대로 인쇄한다 — 실측으로 확인했다.

    python3 tools/bake_proposal.py --docx docs/proposal.docx --out _site/proposal.pdf
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TOC_HEADING = "목차"
TAB_INCHES = 6.3
"""오른쪽 탭 자리. 본문 폭(용지 8.5인치에서 여백 0.95인치 둘을 뺀 값)보다 조금 안쪽이다."""


def render(docx: Path, outdir: Path) -> Path:
    """LibreOffice로 굽는다. **글꼴이 없으면 조용히 네모가 되므로** 호출자가 본문을 확인한다."""
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ("soffice", "--headless", "--convert-to", "pdf", str(docx), "--outdir", str(outdir)),
        check=True,
        capture_output=True,
        timeout=600,
    )
    made = outdir / f"{docx.stem}.pdf"
    if not made.exists():
        raise RuntimeError(f"PDF가 안 나왔다: {made}")
    return made


def page_texts(pdf: Path) -> list[str]:
    """쪽별 글자. `pdftotext`는 쪽 사이에 폼피드를 넣는다."""
    done = subprocess.run(
        ("pdftotext", "-layout", str(pdf), "-"),
        check=True,
        capture_output=True,
        text=True,
        timeout=600,
    )
    return done.stdout.split("\f")


def flatten(text: str) -> str:
    """띄어쓰기를 지운 비교용 꼴. **PDF는 줄바꿈과 공백을 제 마음대로 넣는다.**"""
    return re.sub(r"\s+", "", text)


def locate(pages: list[str], titles: list[str]) -> dict[str, int]:
    """제목마다 **본문에서 처음 나오는 쪽**. 목차 쪽 자신은 건너뛴다.

    목차에도 같은 글자가 있으므로 거기서 세면 전부 목차 쪽 번호가 된다.
    """
    start = 0
    for index, text in enumerate(pages):
        if flatten(TOC_HEADING) in flatten(text):
            start = index + 1
            break
    found: dict[str, int] = {}
    for title in titles:
        needle = flatten(title)
        for index in range(start, len(pages)):
            if needle and needle in flatten(pages[index]):
                found[title] = index + 1
                break
    return found


def toc_paragraphs(document: object, titles: set[str]) -> list[object]:
    """목차 제목 아래에 이어지는 항목 문단들."""
    rows: list[object] = []
    seen_heading = False
    for paragraph in document.paragraphs:  # type: ignore[attr-defined]
        text = paragraph.text.strip()
        if not seen_heading:
            seen_heading = text == TOC_HEADING
            continue
        if text in titles:
            rows.append(paragraph)
        elif rows:
            break
    return rows


def stamp(source: Path, target: Path, numbers: dict[str, int]) -> int:
    """목차 줄 끝에 점선과 쪽 번호를 붙인 사본을 만든다."""
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches

    document = Document(str(source))
    written = 0
    for paragraph in toc_paragraphs(document, set(numbers)):
        page = numbers[paragraph.text.strip()]
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(TAB_INCHES), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        run = paragraph.add_run(f"\t{page}")
        run.font.size = paragraph.runs[0].font.size if paragraph.runs else None
        written += 1
    document.save(str(target))
    return written


def titles_of(source: Path) -> list[str]:
    from docx import Document

    document = Document(str(source))
    seen_heading = False
    rows: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not seen_heading:
            seen_heading = text == TOC_HEADING
            continue
        if not text:
            continue
        if rows and text.startswith("Part") is False and len(rows) > 40:
            break
        if re.match(r"^(Part|\d+\.|□ )", text):
            rows.append(text)
        elif rows:
            break
    return rows


def bake(docx: Path, out: Path) -> int:
    with tempfile.TemporaryDirectory() as box:
        work = Path(box)
        first = render(docx, work / "first")
        titles = titles_of(docx)
        if not titles:
            print("목차 항목을 못 찾았다", file=sys.stderr)
            return 2
        numbers = locate(page_texts(first), titles)
        missing = [title for title in titles if title not in numbers]
        if missing:
            print(f"본문에서 못 찾은 목차 항목 {len(missing)}개: {missing[:3]}", file=sys.stderr)
            return 1

        numbered = work / docx.name
        stamp(docx, numbered, numbers)
        second = render(numbered, work / "second")

        # **다시 센다.** 숫자를 붙여도 쪽이 밀리지 않는다는 것이 이 방식의 전제다.
        again = locate(page_texts(second), titles)
        drifted = {title: (numbers[title], again.get(title)) for title in titles}
        drifted = {name: pair for name, pair in drifted.items() if pair[0] != pair[1]}
        if drifted:
            print(f"쪽이 밀렸다: {list(drifted.items())[:3]}", file=sys.stderr)
            return 1

        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(second, out)
    print(f"구웠다 · 목차 {len(numbers)}줄에 쪽 번호 · {out}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="기획서 PDF 굽기 + 목차 쪽 번호 (D-0246)")
    parser.add_argument("--docx", type=Path, default=ROOT / "docs" / "proposal.docx")
    parser.add_argument("--out", type=Path, default=ROOT / "_site" / "proposal.pdf")
    args = parser.parse_args()
    if not args.docx.exists():
        print(f"기획서가 없다: {args.docx}  `make proposal`", file=sys.stderr)
        return 2
    return bake(args.docx, args.out)


if __name__ == "__main__":
    raise SystemExit(main())
